# -*- coding: utf-8 -*-
"""One description of a schematic, rendered either to matplotlib or to native
PowerPoint shapes.

The manuscript needs a vector figure; the reviewer needs a deck in which every
box, arrow and label is an individual PowerPoint object. Describing the figure
once and emitting it twice is the only way to keep those two in step.

Coordinates are inches with the origin at the bottom left, matching matplotlib.
The PowerPoint emitter flips y.
"""
import numpy as np

import matplotlib.pyplot as plt
from matplotlib.patches import (Rectangle, FancyBboxPatch, Ellipse, Polygon,
                                FancyArrowPatch)

FONT = "Arial"
PT = 1.0 / 72.0


def _rgb(c):
    import matplotlib.colors as mc
    r, g, b = mc.to_rgb(c)
    return int(round(r * 255)), int(round(g * 255)), int(round(b * 255))


class Scene(object):
    def __init__(self, w, h):
        self.w, self.h = w, h
        self.items = []

    # ---------------------------------------------------------- primitives
    def rect(self, x, y, w, h, fill=None, line=None, lw=1.0, dash=None,
             radius=0.0, z=1):
        self.items.append(dict(kind="rect", x=x, y=y, w=w, h=h, fill=fill,
                               line=line, lw=lw, dash=dash, radius=radius, z=z))

    def ellipse(self, cx, cy, rx, ry=None, fill=None, line=None, lw=1.0, z=1):
        ry = rx if ry is None else ry
        self.items.append(dict(kind="ellipse", cx=cx, cy=cy, rx=rx, ry=ry,
                               fill=fill, line=line, lw=lw, z=z))

    def line(self, x0, y0, x1, y1, color="#000000", lw=1.0, dash=None,
             arrow=False, z=2, head=9.0):
        self.items.append(dict(kind="line", x0=x0, y0=y0, x1=x1, y1=y1,
                               color=color, lw=lw, dash=dash, arrow=arrow,
                               head=head, z=z))

    def arrow(self, x0, y0, x1, y1, color="#000000", lw=1.0, dash=None,
              head=9.0, z=2):
        self.line(x0, y0, x1, y1, color, lw, dash, arrow=True, z=z,
                  head=head)

    def poly(self, pts, fill=None, line=None, lw=1.0, z=1):
        self.items.append(dict(kind="poly", pts=list(pts), fill=fill,
                               line=line, lw=lw, z=z))

    def text(self, x, y, s, size=8.0, color="#12293F", bold=False,
             italic=False, ha="left", va="center", width=None, lead=1.30, z=5):
        self.items.append(dict(kind="text", x=x, y=y, s=s, size=size,
                               color=color, bold=bold, italic=italic, ha=ha,
                               va=va, width=width, lead=lead, z=z))

    # ------------------------------------------------------------ matplotlib
    def to_mpl(self):
        fig = plt.figure(figsize=(self.w, self.h))
        ax = fig.add_axes([0, 0, 1, 1])
        ax.set_xlim(0, self.w)
        ax.set_ylim(0, self.h)
        ax.set_xticks([]); ax.set_yticks([])
        for sp in ax.spines.values():
            sp.set_visible(False)
        for it in sorted(self.items, key=lambda d: d["z"]):
            k = it["kind"]
            if k == "rect":
                kw = dict(facecolor=it["fill"] or "none",
                          edgecolor=it["line"] or "none", lw=it["lw"],
                          zorder=it["z"])
                if it["dash"]:
                    kw["linestyle"] = it["dash"]
                if it["radius"]:
                    ax.add_patch(FancyBboxPatch(
                        (it["x"], it["y"]), it["w"], it["h"],
                        boxstyle="round,pad=0,rounding_size=%.3f" % it["radius"],
                        **kw))
                else:
                    ax.add_patch(Rectangle((it["x"], it["y"]), it["w"],
                                           it["h"], **kw))
            elif k == "ellipse":
                ax.add_patch(Ellipse((it["cx"], it["cy"]), 2 * it["rx"],
                                     2 * it["ry"],
                                     facecolor=it["fill"] or "none",
                                     edgecolor=it["line"] or "none",
                                     lw=it["lw"], zorder=it["z"]))
            elif k == "line":
                if it["arrow"]:
                    ax.add_patch(FancyArrowPatch(
                        (it["x0"], it["y0"]), (it["x1"], it["y1"]),
                        arrowstyle="-|>", mutation_scale=it.get("head", 9),
                        lw=it["lw"],
                        color=it["color"], linestyle=it["dash"] or "-",
                        shrinkA=0, shrinkB=0, zorder=it["z"]))
                else:
                    ax.plot([it["x0"], it["x1"]], [it["y0"], it["y1"]],
                            lw=it["lw"], color=it["color"],
                            linestyle=it["dash"] or "-", zorder=it["z"],
                            solid_capstyle="butt")
            elif k == "poly":
                ax.add_patch(Polygon(it["pts"], closed=True,
                                     facecolor=it["fill"] or "none",
                                     edgecolor=it["line"] or "none",
                                     lw=it["lw"], zorder=it["z"]))
            elif k == "text":
                ax.text(it["x"], it["y"], it["s"], fontsize=it["size"],
                        color=it["color"], ha=it["ha"], va=it["va"],
                        fontweight="bold" if it["bold"] else "normal",
                        style="italic" if it["italic"] else "normal",
                        linespacing=it["lead"], zorder=it["z"],
                        fontfamily=FONT)
        return fig

    # --------------------------------------------------------------- pptx
    def to_pptx(self, slide, ox, oy, scale=1.0):
        """Emit every item as its own PowerPoint shape.

        ox, oy are the top-left of the drawing on the slide, in inches.
        """
        from pptx.util import Inches, Pt, Emu
        from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
        from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
        from pptx.dml.color import RGBColor
        from pptx.oxml.ns import qn

        def X(v):
            return Inches(ox + v * scale)

        def Y(v):                                   # flip to top-left origin
            return Inches(oy + (self.h - v) * scale)

        def style(shape, fill, line, lw, dash=None):
            if fill:
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor(*_rgb(fill))
            else:
                shape.fill.background()
            if line:
                shape.line.color.rgb = RGBColor(*_rgb(line))
                shape.line.width = Pt(lw)
                if dash:
                    shape.line._get_or_add_ln().append(
                        _dash_el(qn, "sysDot" if dash == ":" else "sysDash"))
            else:
                shape.line.fill.background()
            shape.shadow.inherit = False

        def _dash_el(qn_, val):
            from pptx.oxml import parse_xml
            from pptx.oxml.ns import nsdecls
            return parse_xml('<a:prstDash %s val="%s"/>' % (nsdecls("a"), val))

        for it in sorted(self.items, key=lambda d: d["z"]):
            k = it["kind"]
            if k == "rect":
                shp = slide.shapes.add_shape(
                    MSO_SHAPE.ROUNDED_RECTANGLE if it["radius"]
                    else MSO_SHAPE.RECTANGLE,
                    X(it["x"]), Y(it["y"] + it["h"]),
                    Inches(it["w"] * scale), Inches(it["h"] * scale))
                if it["radius"]:
                    adj = min(0.5, it["radius"] / min(it["w"], it["h"]))
                    shp.adjustments[0] = adj
                style(shp, it["fill"], it["line"], it["lw"], it["dash"])
                shp.text_frame.text = ""
            elif k == "ellipse":
                shp = slide.shapes.add_shape(
                    MSO_SHAPE.OVAL, X(it["cx"] - it["rx"]),
                    Y(it["cy"] + it["ry"]),
                    Inches(2 * it["rx"] * scale), Inches(2 * it["ry"] * scale))
                style(shp, it["fill"], it["line"], it["lw"])
            elif k == "line":
                cn = slide.shapes.add_connector(
                    MSO_CONNECTOR.STRAIGHT, X(it["x0"]), Y(it["y0"]),
                    X(it["x1"]), Y(it["y1"]))
                cn.line.color.rgb = RGBColor(*_rgb(it["color"]))
                cn.line.width = Pt(it["lw"])
                ln = cn.line._get_or_add_ln()
                if it["dash"]:
                    ln.append(_dash_el(qn, "sysDash"))
                if it["arrow"]:
                    from pptx.oxml import parse_xml
                    from pptx.oxml.ns import nsdecls
                    ln.append(parse_xml(
                        '<a:tailEnd %s type="triangle" w="med" len="med"/>'
                        % nsdecls("a")))
            elif k == "poly":
                pts = it["pts"]
                bld = slide.shapes.build_freeform(X(pts[0][0]), Y(pts[0][1]))
                bld.add_line_segments([(X(px), Y(py)) for px, py in pts[1:]],
                                      close=True)
                shp = bld.convert_to_shape()
                style(shp, it["fill"], it["line"], it["lw"])
            elif k == "text":
                nl = it["s"].count("\n") + 1
                th = nl * it["size"] * PT * it["lead"]
                cw = 0.625 if it["bold"] else 0.545
                tw = it["width"] or (max(len(l) for l in it["s"].split("\n"))
                                     * it["size"] * PT * cw + 0.10)
                bx = {"left": it["x"], "center": it["x"] - tw / 2,
                      "right": it["x"] - tw}[it["ha"]]
                by = {"bottom": it["y"], "center": it["y"] - th / 2,
                      "top": it["y"] - th}[it["va"]]
                tb = slide.shapes.add_textbox(
                    X(bx), Y(by + th), Inches(tw * scale), Inches(th * scale))
                tf = tb.text_frame
                tf.word_wrap = False
                tf.margin_left = tf.margin_right = 0
                tf.margin_top = tf.margin_bottom = 0
                tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                for j, ln_txt in enumerate(it["s"].split("\n")):
                    p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
                    p.alignment = {"left": PP_ALIGN.LEFT,
                                   "center": PP_ALIGN.CENTER,
                                   "right": PP_ALIGN.RIGHT}[it["ha"]]
                    r = p.add_run(); r.text = ln_txt
                    r.font.size = Pt(it["size"])
                    r.font.bold = it["bold"]
                    r.font.italic = it["italic"]
                    r.font.name = FONT
                    r.font.color.rgb = RGBColor(*_rgb(it["color"]))
        return slide
